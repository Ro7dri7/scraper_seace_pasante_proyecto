"""
Enriquecedor de Contactos RNP / OECE (Filtrado por RUC y Directorio)
Consulta la API oficial del Buscador de Proveedores del Estado (RNP / OECE):
https://eap.oece.gob.pe/perfilprov-bus/1.0/ficha/{RUC}

Características:
- Filtrado por RUC: no repite solicitudes y consulta únicamente RUCs únicos.
- Modo Directorio Único (--solo-unicos): Genera un Excel consolidado sin duplicados
  (1 fila por RUC), ideal para prospección comercial y ventas.
- Modo Límite Inteligente (--limit N): Procesa solo N empresas y exporta únicamente
  esas N empresas (sin arrastrar miles de filas vacías).
- Filtrado por RUC específico (--ruc o --rucs).
- Memoria de caché persistente ('cache_contactos_rnp.json'): RUC consultado jamás
  vuelve a pedir petición HTTP (0.0001s desde disco).

"""

import os
import ssl
import csv
import json
import time
import argparse
import datetime
import urllib.request
import concurrent.futures
from collections import defaultdict

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

SSL_CONTEXT = ssl.create_default_context()
SSL_CONTEXT.check_hostname = False
SSL_CONTEXT.verify_mode = ssl.CERT_NONE

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120.0.0.0 Safari/537.36',
    'Accept': 'application/json',
    'Content-Type': 'application/json',
    'Origin': 'https://apps.oece.gob.pe'
}

CACHE_FILE = "cache_contactos_rnp.json"


def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def save_cache(cache):
    try:
        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(cache, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


def fetch_proveedor_ficha(ruc):
    """Consulta la API de la ficha del proveedor en OECE/RNP."""
    clean_r = str(ruc).replace("PE-RUC-", "").strip()
    if len(clean_r) != 11 or not clean_r.isdigit():
        return clean_r, {
            "telefono": "No registra",
            "email": "No registra",
            "apto": "No",
            "habilitado": "No",
            "ficha_url": ""
        }

    url = f"https://eap.oece.gob.pe/perfilprov-bus/1.0/ficha/{clean_r}"
    req = urllib.request.Request(url, headers=HEADERS)
    try:
        with urllib.request.urlopen(req, context=SSL_CONTEXT, timeout=8) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            prov = data.get("proveedorT01", {})
            tels = prov.get("telefonos") or []
            emails = prov.get("emails") or []
            es_apto = prov.get("esAptoContratar", False)
            es_hab = prov.get("esHabilitado", False)

            tel_str = " / ".join(t.strip() for t in tels if t.strip()) if tels else "No registra"
            email_str = " / ".join(e.strip() for e in emails if e.strip()) if emails else "No registra"

            return clean_r, {
                "telefono": tel_str,
                "email": email_str,
                "apto": "SÍ" if es_apto else "No",
                "habilitado": "SÍ" if es_hab else "No",
                "ficha_url": f"https://apps.oece.gob.pe/perfilprov-ui/ficha/{clean_r}"
            }
    except Exception:
        return clean_r, {
            "telefono": "No registra",
            "email": "No registra",
            "apto": "No",
            "habilitado": "No",
            "ficha_url": f"https://apps.oece.gob.pe/perfilprov-ui/ficha/{clean_r}"
        }


def enrich_file(input_file, limit=None, workers=25, solo_unicos=False, filter_rucs=None):
    print("=" * 70)
    print(" ENRIQUECEDOR OFICIAL RNP / OECE (CON FILTRO ESTRICTO POR RUC)")
    print("=" * 70)
    print(f" Archivo de entrada : {input_file}")
    print(f" Modo Directorio    : {'ACTIVADO (1 fila por empresa única)' if solo_unicos else 'Detallado (1 fila por licitación)'}")
    print(f" Workers paralelos  : {workers}")
    print("=" * 70 + "\n")

    if not os.path.exists(input_file):
        print(f"[!] Error: No se encontró el archivo '{input_file}'")
        return

    # 1. Cargar archivo (Soporta tanto CSV como XLSX directamente)
    leads = []
    if input_file.lower().endswith(".xlsx"):
        print(f"[*] Leyendo archivo Excel con openpyxl: {input_file}...")
        wb = openpyxl.load_workbook(input_file, read_only=True)
        ws = wb.active
        row_iter = ws.iter_rows(values_only=True)
        headers = [str(c) if c is not None else "" for c in next(row_iter)]
        for r_vals in row_iter:
            row_dict = {headers[i]: (str(r_vals[i]) if i < len(r_vals) and r_vals[i] is not None else "") for i in range(len(headers))}
            leads.append(row_dict)
    else:
        with open(input_file, mode="r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                leads.append(row)

    print(f"[+] Total de filas en archivo: {len(leads)}")

    # 2. Agrupar por RUC único y recopilar estadísticas de postulación
    ruc_stats = defaultdict(lambda: {
        "razon_social": "",
        "total_postulaciones": 0,
        "total_ganadas": 0,
        "licitaciones": [],
        "monto_total_participado": 0.0
    })

    for l in leads:
        r = l.get("RUC", "").strip()
        if r and len(r) == 11 and r.isdigit():
            name = l.get("Razón Social", "")
            cond = l.get("Condición", "")
            lic = l.get("Licitación", "")
            monto_val = l.get("Monto Referencial (S/.)", 0)
            try:
                monto_float = float(monto_val) if monto_val else 0.0
            except ValueError:
                monto_float = 0.0

            if not ruc_stats[r]["razon_social"]:
                ruc_stats[r]["razon_social"] = name
            ruc_stats[r]["total_postulaciones"] += 1
            if "Ganador" in cond:
                ruc_stats[r]["total_ganadas"] += 1
            if lic and lic not in ruc_stats[r]["licitaciones"]:
                ruc_stats[r]["licitaciones"].append(lic)
            ruc_stats[r]["monto_total_participado"] += monto_float

    unique_rucs = list(ruc_stats.keys())
    print(f"[+] Empresas únicas detectadas (RUCs válidos): {len(unique_rucs)}")

    # Filtrar por lista específica de RUCs si se solicitó
    if filter_rucs:
        target_set = set(filter_rucs)
        unique_rucs = [r for r in unique_rucs if r in target_set]
        print(f"[*] Filtrando exclusivamente por {len(unique_rucs)} RUCs especificados...")

    if limit and limit > 0:
        unique_rucs = unique_rucs[:limit]
        print(f"[*] Aplicando límite: procesando las primeras {len(unique_rucs)} empresas únicas...")

    target_ruc_set = set(unique_rucs)

    # 3. Consultar caché o API en paralelo
    cache = load_cache()
    pending = [r for r in unique_rucs if r not in cache]
    cached_count = len(unique_rucs) - len(pending)
    print(f"\n[*] Estado de Caché Local:")
    print(f"    -> Ya disponibles en disco (0s de espera): {cached_count} empresas")
    print(f"    -> Nuevas a consultar en API de RNP     : {len(pending)} empresas")

    t0 = time.time()
    done_count = 0
    phones_found = sum(1 for r in unique_rucs if r in cache and cache[r]["telefono"] != "No registra")

    if pending:
        with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as executor:
            future_map = {executor.submit(fetch_proveedor_ficha, r): r for r in pending}
            for f in concurrent.futures.as_completed(future_map):
                done_count += 1
                try:
                    ruc, info = f.result()
                    cache[ruc] = info
                    if info["telefono"] != "No registra":
                        phones_found += 1
                    if done_count % 50 == 0 or done_count == len(pending):
                        pct = (done_count / len(pending)) * 100
                        print(f"    [*] Progreso RNP: {done_count}/{len(pending)} ({pct:.1f}%) | Teléfonos hallados: {phones_found}", end="\r")
                    if done_count % 150 == 0:
                        save_cache(cache)
                except Exception:
                    pass
        save_cache(cache)
        print()

    elapsed = time.time() - t0
    total_with_phone = sum(1 for r in unique_rucs if r in cache and cache[r]["telefono"] != "No registra")
    total_with_email = sum(1 for r in unique_rucs if r in cache and cache[r]["email"] != "No registra")

    print(f"\n[+] Proceso completado en {elapsed:.2f} segundos")
    print(f"[+] Empresas con Teléfono Oficial : {total_with_phone}/{len(unique_rucs)} ({(total_with_phone/len(unique_rucs))*100:.1f}%)")
    print(f"[+] Empresas con Email Corporativo: {total_with_email}/{len(unique_rucs)} ({(total_with_email/len(unique_rucs))*100:.1f}%)")

    # 4. Preparar salida según el modo (Directorio Único vs Detallado)
    base_name = os.path.splitext(input_file)[0]

    if solo_unicos:
        # MODO DIRECTORIO: 1 fila por RUC único
        print("\n[*] Generando Directorio Único de Empresas (sin duplicados)...")
        export_rows = []
        for r in unique_rucs:
            st = ruc_stats[r]
            c = cache.get(r, {
                "telefono": "No registra",
                "email": "No registra",
                "habilitado": "No",
                "apto": "No",
                "ficha_url": f"https://apps.oece.gob.pe/perfilprov-ui/ficha/{r}"
            })
            post = st["total_postulaciones"]
            gan = st["total_ganadas"]
            tasa = f"{(gan/post)*100:.1f}%" if post > 0 else "0.0%"
            sample_lics = " | ".join(st["licitaciones"][:3])

            row = {
                "RUC": r,
                "Razón Social": st["razon_social"],
                "Teléfono / Celular (RNP)": c["telefono"],
                "Email de Contacto (RNP)": c["email"],
                "Habilitado RNP": c["habilitado"],
                "Apto para Contratar": c["apto"],
                "Total Postulaciones": post,
                "Licitaciones Ganadas": gan,
                "Tasa de Éxito": tasa,
                "Monto Total Concursado (S/.)": st["monto_total_participado"],
                "Ejemplos de Licitaciones": sample_lics,
                "Ficha RNP (Web)": c["ficha_url"]
            }
            export_rows.append(row)

        out_csv = f"{base_name}_DIRECTORIO_RUC.csv"
        out_xlsx = f"{base_name}_DIRECTORIO_RUC.xlsx"
        headers = list(export_rows[0].keys())

    else:
        # MODO DETALLADO: Mantiene licitaciones, pero filtra solo a los RUCs procesados
        print("\n[*] Asignando datos a las filas detalladas...")
        export_rows = []
        for l in leads:
            r = l.get("RUC", "").strip()
            if r in target_ruc_set:
                c = cache.get(r, {})
                l["Teléfono / Celular (RNP)"] = c.get("telefono", "No registra")
                l["Email de Contacto (RNP)"] = c.get("email", "No registra")
                l["Habilitado RNP"] = c.get("habilitado", "No")
                l["Ficha RNP (Web)"] = c.get("ficha_url", f"https://apps.oece.gob.pe/perfilprov-ui/ficha/{r}")
                export_rows.append(l)

        out_csv = f"{base_name}_ENRIQUECIDO.csv"
        out_xlsx = f"{base_name}_ENRIQUECIDO.xlsx"
        
        raw_headers = list(export_rows[0].keys())
        priority = ["RUC", "Razón Social", "Condición", "Teléfono / Celular (RNP)", "Email de Contacto (RNP)", "Habilitado RNP", "Ficha RNP (Web)"]
        headers = priority + [h for h in raw_headers if h not in priority]

    # Guardar CSV
    with open(out_csv, mode="w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(export_rows)
    print(f"[+] Archivo CSV generado: {out_csv} ({len(export_rows)} filas)")

    # Guardar Excel
    if HAS_OPENPYXL:
        export_excel_custom(export_rows, headers, out_xlsx, is_directory=solo_unicos)


def export_excel_custom(rows, headers, filename, is_directory=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Directorio RNP" if is_directory else "Prospectos RNP"
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_num)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, r_dict in enumerate(rows, 2):
        for col_idx, col_name in enumerate(headers, 1):
            val = r_dict.get(col_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_thin
            cell.font = Font(name="Calibri", size=10)

            if col_name in ["Teléfono / Celular (RNP)", "Email de Contacto (RNP)"]:
                if val and val != "No registra":
                    cell.font = Font(name="Calibri", size=10, bold=True, color="1F4E78")

            elif col_name in ["Monto Referencial (S/.)", "Monto Total Concursado (S/.)"] and isinstance(val, (int, float)):
                cell.number_format = '"S/." #,##0.00'

            elif col_name == "Ficha RNP (Web)" and str(val).startswith("http"):
                cell.value = f'=HYPERLINK("{val}", "Ver Ficha RNP")'
                cell.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
                cell.alignment = Alignment(horizontal="center")

            elif col_name == "Ver Proceso (Portal OECE)" and str(val).startswith("http"):
                cell.value = f'=HYPERLINK("{val}", "Ver en OECE")'
                cell.font = Font(name="Calibri", size=10, color="0563C1", underline="single")
                cell.alignment = Alignment(horizontal="center")

            elif col_name in ["RUC", "Habilitado RNP", "Apto para Contratar", "Total Postulaciones", "Licitaciones Ganadas", "Tasa de Éxito", "Condición"]:
                cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col[:150]:
            val_str = str(cell.value or '')
            if val_str.startswith("=HYPERLINK"):
                val_str = "Ver Ficha RNP"
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    try:
        wb.save(filename)
        print(f"[+] Archivo Excel generado con éxito: {filename} ({len(rows)} filas)")
    except PermissionError:
        fallback = filename.replace(".xlsx", f"_{int(time.time()) % 10000}.xlsx")
        wb.save(fallback)
        print(f"[!] Archivo abierto. Se guardó copia como: {fallback}")


def main():
    parser = argparse.ArgumentParser(description="Enriquecedor oficial RNP / OECE con filtro estricto por RUC")
    parser.add_argument("--input", type=str, required=True, help="Archivo CSV de prospectos")
    parser.add_argument("--limit", type=int, default=None, help="Límite máximo de empresas únicas a consultar")
    parser.add_argument("--solo-unicos", action="store_true", help="Generar directorio único (1 fila por RUC, sin duplicados)")
    parser.add_argument("--ruc", type=str, default=None, help="Filtrar por un RUC específico (ej. 20611555564)")
    parser.add_argument("--rucs", type=str, default=None, help="Filtrar por lista de RUCs separados por coma")
    parser.add_argument("--workers", type=int, default=25, help="Hilos de consulta paralelos (default: 25)")

    args = parser.parse_args()

    filter_rucs = None
    if args.ruc:
        filter_rucs = [args.ruc.strip()]
    elif args.rucs:
        filter_rucs = [r.strip() for r in args.rucs.split(",") if r.strip()]

    enrich_file(args.input, limit=args.limit, workers=args.workers, solo_unicos=args.solo_unicos, filter_rucs=filter_rucs)


if __name__ == "__main__":
    main()
